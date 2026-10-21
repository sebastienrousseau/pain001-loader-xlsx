# SPDX-FileCopyrightText: 2026 Sebastien Rousseau <sebastian.rousseau@gmail.com>
# SPDX-License-Identifier: Apache-2.0 OR MIT
"""Cells starting with =, +, - or @ reach the payment as written.

0.0.72 prefixed such cells with a single quote, as the core's CSV
loader then did (CWE-1236). pain001 writes ISO 20022 XML, which no
spreadsheet evaluates, so the quote protected nothing and changed the
payment: ``-5`` became ``'-5`` and ``+44 20 7946 0000`` gained a
quote. pain001 0.0.74 preserves raw CSV values by default (its
ADR-0008); this loader does the same, so a sheet and its CSV export
still render the same XML.
"""

from __future__ import annotations

import csv
import pathlib
from importlib import metadata

import openpyxl
import pytest
from packaging.version import Version
from pain001.constants import TEMPLATES_DIR
from pain001.data.loader import load_payment_data
from pain001.plugins.registry import registry
from pain001.xml.generate_xml import generate_xml_string

from pain001_loader_xlsx._normalise import to_text
from pain001_loader_xlsx.loader import XlsxLoader

RAW = [
    "+44 20 7946 0000",
    "-INV-0042",
    "=A1",
    "@x",
    "\t=calc",
]


@pytest.mark.parametrize("text", RAW)
def test_to_text_keeps_the_text(text: str) -> None:
    """No quote is added in front of a formula-like string."""
    assert to_text(text, None) == text


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [(-5, "General", "-5"), (-5.0, "General", "-5"), (-1.5, "0.00", "-1.50")],
)
def test_negative_numbers_keep_their_sign(
    value: float, fmt: str, expected: str
) -> None:
    """A negative amount is a number, not a formula."""
    assert to_text(value, fmt) == expected


def test_the_loader_returns_raw_cells(tmp_path: pathlib.Path) -> None:
    """Through the loader: text cells (never formulas) and a number."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["n", "text"])
    for row, text in enumerate(RAW, start=2):
        sheet.cell(row=row, column=1, value=-5)
        cell = sheet.cell(row=row, column=2, value=text)
        cell.data_type = "s"
    target = tmp_path / "raw.xlsx"
    workbook.save(target)

    rows = XlsxLoader().load(str(target)).rows

    assert [r["text"] for r in rows] == RAW
    assert {r["n"] for r in rows} == {"-5"}


@pytest.mark.skipif(
    Version(metadata.version("pain001")) < Version("0.0.74"),
    reason="pain001 before 0.0.74 quotes these cells in its CSV loader",
)
def test_formula_like_cells_render_as_the_csv_does(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Same rows and same XML bytes as the equivalent CSV, on 0.0.74."""
    assets = TEMPLATES_DIR / "pain.001.001.03"
    with (assets / "template.csv").open(newline="", encoding="utf-8") as fh:
        records = list(csv.reader(fh))
    column = records[0].index("remittance_information")
    records[1][column] = "-INV-0042 +44 20 7946 0000"
    workbook = openpyxl.Workbook()
    for record in records:
        workbook.active.append(record)
    for cell in workbook.active[2]:
        cell.data_type = "s"
    xlsx = tmp_path / "payments.xlsx"
    workbook.save(xlsx)
    equivalent = tmp_path / "payments.csv"
    with equivalent.open("w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(records)
    monkeypatch.chdir(tmp_path)
    registry.reset()

    rows_csv = load_payment_data(str(equivalent))
    rows_xlsx = load_payment_data(str(xlsx))

    assert rows_xlsx == rows_csv
    assert rows_xlsx[0]["remittance_information"] == records[1][column]
    args = (
        "pain.001.001.03",
        str(assets / "template.xml"),
        str(assets / "pain.001.001.03.xsd"),
    )
    assert generate_xml_string(rows_xlsx, *args) == generate_xml_string(
        rows_csv, *args
    )
